from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from publish_drafts import PublishDraft, new_publish_draft_id
from publish_platforms import PLATFORMS, classify_media, normalize_tags


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm"}
MEDIA_EXTENSIONS = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS
MAX_XIAOHONGSHU_IMAGES = 18


def _natural_key(path: Path) -> list[object]:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", path.name)
    ]


def _resolve_path(value: str, base_dir: Path) -> Path:
    path = Path(value.strip().strip('"').strip("'"))
    if not path.is_absolute():
        path = base_dir / path
    return path.resolve(strict=False)


def _split_path_values(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    parts = [part.strip() for part in re.split(r"[\n;|]+", text) if part.strip()]
    if len(parts) == 1 and "," in parts[0]:
        parts = [part.strip() for part in parts[0].split(",") if part.strip()]
    return parts


def _direct_media_files(directory: Path) -> list[Path]:
    if not directory.is_dir():
        return []
    return sorted(
        (
            path
            for path in directory.iterdir()
            if path.is_file()
            and not path.name.startswith(".")
            and path.suffix.lower() in MEDIA_EXTENSIONS
        ),
        key=_natural_key,
    )


def _normalize_media_type(value: str, media_paths: list[Path]) -> str:
    normalized = value.strip().lower()
    if normalized in {"图文", "图片", "image", "images", "笔记"}:
        return "image"
    if normalized in {"视频", "video", "videos"}:
        return "video"
    videos = [path for path in media_paths if path.suffix.lower() in VIDEO_EXTENSIONS]
    images = [path for path in media_paths if path.suffix.lower() in IMAGE_EXTENSIONS]
    if videos and not images:
        return "video"
    return "video" if len(videos) == 1 else "image"


def _make_draft(
    *,
    media_paths: list[Path],
    media_type: str,
    title: str,
    description: str,
    tags: list[str],
    cover_path: Path | None = None,
    scheduled_at: str = "",
) -> PublishDraft:
    return PublishDraft(
        draft_id=new_publish_draft_id(),
        source_page="批量导入",
        media_paths=[str(path) for path in media_paths],
        media_type=media_type,
        title=title.strip(),
        description=description.strip(),
        tags=list(tags),
        cover_path=str(cover_path) if cover_path else "",
        scheduled_at=scheduled_at.strip(),
    )


def _drafts_from_directory(
    directory: Path,
    *,
    title: str = "",
    description: str = "",
    tags: list[str] | None = None,
    media_type: str = "",
    cover_path: Path | None = None,
    scheduled_at: str = "",
) -> list[PublishDraft]:
    files = _direct_media_files(directory)
    images = [path for path in files if path.suffix.lower() in IMAGE_EXTENSIONS]
    videos = [path for path in files if path.suffix.lower() in VIDEO_EXTENSIONS]
    normalized_type = _normalize_media_type(media_type, files)
    draft_title = title.strip() or directory.name
    draft_tags = list(tags or [])
    drafts: list[PublishDraft] = []

    if normalized_type == "video":
        selected_cover = cover_path or (images[0] if images else None)
        for index, video in enumerate(videos):
            video_cover = selected_cover
            if not cover_path and images:
                video_cover = images[min(index, len(images) - 1)]
            drafts.append(
                _make_draft(
                    media_paths=[video],
                    media_type="video",
                    title=draft_title,
                    description=description,
                    tags=draft_tags,
                    cover_path=video_cover,
                    scheduled_at=scheduled_at,
                )
            )
        return drafts

    selected_images = list(images)
    if cover_path and cover_path not in selected_images:
        selected_images.insert(0, cover_path)
    if selected_images:
        drafts.append(
            _make_draft(
                media_paths=selected_images,
                media_type="image",
                title=draft_title,
                description=description,
                tags=draft_tags,
                cover_path=cover_path or selected_images[0],
                scheduled_at=scheduled_at,
            )
        )
    return drafts


def discover_publish_drafts(root: str | Path) -> list[PublishDraft]:
    root_path = Path(root).resolve()
    if root_path.is_file():
        media_type = (
            "video"
            if root_path.suffix.lower() in VIDEO_EXTENSIONS
            else "image"
        )
        return [
            _make_draft(
                media_paths=[root_path],
                media_type=media_type,
                title=root_path.stem,
                description="",
                tags=[],
                cover_path=root_path if media_type == "image" else None,
            )
        ]
    if not root_path.is_dir():
        raise ValueError(f"发布素材目录不存在：{root_path}")

    media = [
        path
        for path in root_path.rglob("*")
        if path.is_file()
        and not path.name.startswith(".")
        and path.suffix.lower() in MEDIA_EXTENSIONS
    ]
    if not media:
        return []

    parent_dirs = sorted({path.parent for path in media}, key=lambda path: str(path).lower())
    drafts: list[PublishDraft] = []
    for directory in parent_dirs:
        drafts.extend(_drafts_from_directory(directory))
    return drafts


def _row_value(row: dict[str, object], aliases: tuple[str, ...]) -> str:
    normalized = {
        str(key).strip().lower(): value
        for key, value in row.items()
        if key is not None
    }
    for alias in aliases:
        value = normalized.get(alias.strip().lower())
        if value is not None:
            return str(value).strip()
    return ""


def _rows_from_excel(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [
            str(value).strip() if value is not None else ""
            for value in rows[0]
        ]
        result: list[dict[str, object]] = []
        for values in rows[1:]:
            if not any(value not in (None, "") for value in values):
                continue
            result.append(
                {
                    headers[index]: values[index]
                    for index in range(min(len(headers), len(values)))
                }
            )
        return result
    finally:
        workbook.close()


def load_publish_drafts_from_excel(
    path: str | Path,
    *,
    base_dir: str | Path | None = None,
) -> list[PublishDraft]:
    excel_path = Path(path).resolve()
    if excel_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("批量发布仅支持 .xlsx 或 .xlsm 文件。")
    root = Path(base_dir).resolve() if base_dir else excel_path.parent
    rows = _rows_from_excel(excel_path)
    drafts: list[PublishDraft] = []

    for row in rows:
        title = _row_value(row, ("标题", "视频主题", "主题", "title"))
        description = _row_value(
            row,
            ("正文", "简介", "文案", "description", "内容"),
        )
        raw_tags = _row_value(row, ("标签", "标签关键词", "关键词", "tags"))
        type_text = _row_value(row, ("类型", "发布类型", "type"))
        scheduled_at = _row_value(row, ("计划发布时间", "定时发布时间", "schedule_at"))
        raw_media = _row_value(row, ("素材路径", "图片路径", "视频路径", "media", "media_path"))
        raw_cover = _row_value(row, ("封面路径", "封面", "cover", "cover_path"))
        cover_path = _resolve_path(raw_cover, root) if raw_cover else None

        media_values = _split_path_values(raw_media)
        resolved_media = [_resolve_path(value, root) for value in media_values]
        expanded_paths: list[Path] = []
        row_drafts: list[PublishDraft] = []
        normalized_type = _normalize_media_type(type_text, resolved_media)
        for media_path in resolved_media:
            if media_path.is_dir():
                row_drafts.extend(
                    _drafts_from_directory(
                        media_path,
                        title=title,
                        description=description,
                        tags=normalize_tags(raw_tags, 10),
                        media_type=type_text,
                        cover_path=cover_path,
                        scheduled_at=scheduled_at,
                    )
                )
            else:
                expanded_paths.append(media_path)

        if expanded_paths:
            normalized_type = _normalize_media_type(type_text, expanded_paths)
            if normalized_type == "video":
                selected_media = [
                    path
                    for path in expanded_paths
                    if path.suffix.lower() in VIDEO_EXTENSIONS
                ][:1]
            else:
                selected_media = [
                    path
                    for path in expanded_paths
                    if path.suffix.lower() in IMAGE_EXTENSIONS
                ]
                if cover_path and cover_path not in selected_media:
                    selected_media.insert(0, cover_path)
            row_drafts.append(
                _make_draft(
                    media_paths=selected_media,
                    media_type=normalized_type,
                    title=title or (selected_media[0].stem if selected_media else ""),
                    description=description,
                    tags=normalize_tags(raw_tags, 10),
                    cover_path=cover_path
                    or (selected_media[0] if normalized_type == "image" and selected_media else None),
                    scheduled_at=scheduled_at,
                )
            )
        elif not row_drafts:
            row_drafts.append(
                _make_draft(
                    media_paths=[],
                    media_type=normalized_type,
                    title=title,
                    description=description,
                    tags=normalize_tags(raw_tags, 10),
                    cover_path=cover_path,
                    scheduled_at=scheduled_at,
                )
            )

        drafts.extend(row_drafts)
    return [validate_publish_draft(draft) for draft in drafts]


def build_publish_fingerprint(draft: PublishDraft) -> str:
    digest = hashlib.sha256()
    for raw_path in sorted(draft.media_paths):
        path = Path(raw_path)
        try:
            stat = path.stat()
            identity = f"{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}"
        except OSError:
            identity = str(path.resolve(strict=False))
        digest.update(identity.encode("utf-8", errors="replace"))
    digest.update(draft.title.strip().encode("utf-8", errors="replace"))
    digest.update(draft.description.strip().encode("utf-8", errors="replace"))
    return digest.hexdigest()


def validate_publish_draft(
    draft: PublishDraft,
    *,
    platform_key: str = "xiaohongshu",
) -> PublishDraft:
    profile = PLATFORMS.get(platform_key, PLATFORMS["xiaohongshu"])
    errors: list[str] = []
    warnings: list[str] = []
    media = [Path(path) for path in draft.media_paths]
    missing = [path for path in media if not path.exists()]
    if not media:
        errors.append("没有选择图片或视频素材。")
    if missing:
        errors.append(
            "素材文件不存在：" + "、".join(path.name for path in missing[:5])
        )

    images, videos, unsupported = classify_media(media)
    if unsupported:
        errors.append(
            "存在不支持的文件：" + "、".join(path.name for path in unsupported[:5])
        )
    if draft.media_type == "video":
        if len(videos) != 1:
            errors.append("视频发布包必须且只能包含 1 个视频。")
    else:
        if not images:
            errors.append("图文发布包至少需要 1 张图片。")
        if videos:
            errors.append("图文发布包中不能混入视频文件。")
        if len(images) > MAX_XIAOHONGSHU_IMAGES:
            errors.append(
                f"图文发布包包含 {len(images)} 张图片，"
                f"超过当前建议上限 {MAX_XIAOHONGSHU_IMAGES} 张。"
            )

    title = draft.title.strip()
    description = draft.description.strip()
    if not title:
        errors.append("标题不能为空。")
    elif len(title) > profile.title_limit:
        errors.append(
            f"标题 {len(title)} 字，超过{profile.name}建议上限 "
            f"{profile.title_limit} 字。"
        )
    if len(description) > profile.description_limit:
        warnings.append(
            f"正文 {len(description)} 字，超过{profile.name}建议上限 "
            f"{profile.description_limit} 字，发布时将被截断。"
        )
    if not draft.tags:
        warnings.append("尚未填写标签。")
    elif len(draft.tags) > profile.max_tags:
        warnings.append(
            f"标签 {len(draft.tags)} 个，超过建议上限 "
            f"{profile.max_tags} 个，只保留前 {profile.max_tags} 个。"
        )
    if draft.scheduled_at:
        try:
            schedule = datetime.strptime(draft.scheduled_at, "%Y-%m-%d %H:%M")
            if schedule < datetime.now():
                warnings.append("计划发布时间早于当前时间。")
        except ValueError:
            errors.append("计划发布时间格式应为 YYYY-MM-DD HH:MM。")

    draft.validation_errors = errors
    draft.validation_warnings = warnings
    draft.fingerprint = build_publish_fingerprint(draft)
    draft.status = "已校验" if not errors else "需修正"
    return draft


def validate_publish_drafts(
    drafts: list[PublishDraft],
    *,
    platform_key: str = "xiaohongshu",
) -> list[PublishDraft]:
    return [
        validate_publish_draft(draft, platform_key=platform_key)
        for draft in drafts
    ]


def dedupe_publish_drafts(
    drafts: list[PublishDraft],
    *,
    existing: list[PublishDraft] | None = None,
) -> tuple[list[PublishDraft], int]:
    seen = {
        draft.fingerprint or build_publish_fingerprint(draft)
        for draft in (existing or [])
    }
    unique: list[PublishDraft] = []
    duplicates = 0
    for draft in drafts:
        fingerprint = draft.fingerprint or build_publish_fingerprint(draft)
        if fingerprint in seen:
            duplicates += 1
            continue
        seen.add(fingerprint)
        unique.append(draft)
    return unique, duplicates


def create_publish_batch_template(path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "批量发布"
    sheet.append(
        [
            "类型",
            "素材路径",
            "封面路径",
            "标题",
            "正文",
            "标签",
            "计划发布时间",
        ]
    )
    sheet.append(
        [
            "图文",
            "D:/素材/笔记01",
            "D:/素材/笔记01/01_封面.png",
            "久坐党别忽略这件事",
            "1、减少长时间久坐\n2、如厕不刷手机\n3、做好轻柔清洁",
            "肛周护理 久坐党 健康科普",
            "2026-10-01 10:00",
        ]
    )
    sheet.append(
        [
            "视频",
            "D:/素材/视频01.mp4",
            "D:/素材/视频01_封面.jpg",
            "三个日常护理习惯",
            "日常养护从减少反复摩擦开始。",
            "肛周护理 健康生活",
            "2026-10-01 20:00",
        ]
    )
    for column, width in {
        "A": 12,
        "B": 42,
        "C": 42,
        "D": 30,
        "E": 56,
        "F": 30,
        "G": 20,
    }.items():
        sheet.column_dimensions[column].width = width
    workbook.save(target)
    return target
