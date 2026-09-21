from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from prompts import build_image_prompt


@dataclass(slots=True)
class BatchItem:
    index: int
    theme: str
    copy: str
    tags: str = ""
    style: str = ""
    image_prompts: list[str] = field(default_factory=list)
    video_prompt: str = ""
    image_paths: list[Path] = field(default_factory=list)
    image_urls: list[str] = field(default_factory=list)
    video_paths: list[Path] = field(default_factory=list)
    status: str = "待处理"


def _pick(row: dict[str, object], names: tuple[str, ...]) -> str:
    normalized = {str(key).strip().lower(): value for key, value in row.items() if key is not None}
    for name in names:
        value = normalized.get(name.strip().lower())
        if value is not None:
            return str(value).strip()
    return ""


def _split_numbered_prompts(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    pattern = re.compile(r"(图\s*\d+\s*[：:])")
    parts = pattern.split(text)
    if len(parts) >= 3:
        common = parts[0].strip()
        prompts: list[str] = []
        for index in range(1, len(parts), 2):
            marker = parts[index].strip()
            content = parts[index + 1].strip() if index + 1 < len(parts) else ""
            if content:
                prompts.append(f"{common} {marker}{content}".strip())
        if prompts:
            return prompts

    lines = [line.strip(" -\t") for line in text.splitlines() if line.strip(" -\t")]
    if len(lines) > 1:
        return lines
    return [text]


def _rows_from_csv(path: Path) -> list[dict[str, object]]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别 CSV 文件编码。")


def _rows_from_xlsx(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        result: list[dict[str, object]] = []
        for values in rows[1:]:
            if not any(value not in (None, "") for value in values):
                continue
            result.append({headers[index]: values[index] for index in range(min(len(headers), len(values)))})
        return result
    finally:
        workbook.close()


def load_batch_items(
    path: str | Path,
    *,
    default_style: str,
    max_prompts_per_row: int = 6,
) -> list[BatchItem]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        rows = _rows_from_csv(file_path)
    elif file_path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _rows_from_xlsx(file_path)
    else:
        raise ValueError("仅支持 .xlsx、.xlsm 和 .csv 文件。")

    items: list[BatchItem] = []
    for offset, row in enumerate(rows, start=1):
        theme = _pick(row, ("视频主题", "主题", "标题", "title"))
        copy = _pick(row, ("简介", "每张图文案", "图片文案", "文案", "copy", "description"))
        tags = _pick(row, ("标签关键词", "标签", "关键词", "keywords", "tags"))
        style = _pick(row, ("图片风格", "风格", "style")) or default_style
        prompt_text = _pick(row, ("生图提示词", "图像提示词", "图片提示词", "prompt", "image_prompt"))
        video_prompt = _pick(row, ("视频提示词", "视频prompt", "video_prompt"))

        if not theme and not copy and not prompt_text:
            continue

        image_prompts = _split_numbered_prompts(prompt_text)[:max_prompts_per_row]
        if not image_prompts:
            image_prompts = [build_image_prompt(theme, copy, style)]

        items.append(
            BatchItem(
                index=offset,
                theme=theme or copy[:30] or f"主题{offset}",
                copy=copy,
                tags=tags,
                style=style,
                image_prompts=image_prompts,
                video_prompt=video_prompt,
            )
        )
    return items


def create_batch_template(path: str | Path) -> Path:
    target = Path(path)
    if target.suffix.lower() == ".csv":
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["主题", "每张图文案", "风格", "生图提示词", "视频提示词", "标签关键词"])
            writer.writerow(
                [
                    "久坐党如厕习惯",
                    "用四个步骤说明如厕时不要长时间刷手机",
                    "清新治愈",
                    "图1：标题海报；图2：手机留在门外；图3：结束点时钟；图4：轻柔清洁清单",
                    "镜头缓慢推进，纸片和图标依次出现",
                    "#肛周护理 #久坐党 #健康科普",
                ]
            )
        return target

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "批量文案"
    sheet.append(["主题", "每张图文案", "风格", "生图提示词", "视频提示词", "标签关键词"])
    sheet.append(
        [
            "久坐党如厕习惯",
            "用四个步骤说明如厕时不要长时间刷手机",
            "清新治愈",
            "图1：标题海报；图2：手机留在门外；图3：结束点时钟；图4：轻柔清洁清单",
            "镜头缓慢推进，纸片和图标依次出现",
            "#肛周护理 #久坐党 #健康科普",
        ]
    )
    for column, width in {"A": 26, "B": 42, "C": 14, "D": 60, "E": 42, "F": 30}.items():
        sheet.column_dimensions[column].width = width
    workbook.save(target)
    return target
